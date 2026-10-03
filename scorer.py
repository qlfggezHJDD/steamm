"""
Steam Dead Game Ranker — Render + Discord edition
Scoring: log(reviews) 50% + solo dev 30% + price 20%
"""

import json
import math
import time
import re
import os
import io
import datetime
import requests
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

# ── CONFIG ────────────────────────────────────────────────────────────────────
DISCORD_WEBHOOK = os.environ["DISCORD_WEBHOOK_URL"]
LINKS_FILE      = os.environ.get("LINKS_FILE", "steam_links.txt")
TOP_N           = int(os.environ.get("TOP_N", "30"))
BATCH_DELAY_S   = float(os.environ.get("BATCH_DELAY_S", "1.5"))
MAX_PER_RUN     = int(os.environ.get("MAX_PER_RUN", "9999"))
PROGRESS_FILE   = "/tmp/steam_progress.json"

# ── SCORE ─────────────────────────────────────────────────────────────────────

def score_game(reviews: int, dev_games: int, price: float) -> float:
    # 50% — review count (log scale: 0 reviews = 100, 10k+ reviews ≈ 0)
    review_score = max(0, 100 - 33 * math.log10(reviews + 1))

    # 30% — dev game count (1 game = 100, 10+ games = 0)
    dev_score = max(0, 100 - (dev_games - 1) * 11.1)

    # 20% — price ($0–5 = 100, $5–15 = 60, $15–30 = 20, $30+ = 0)
    if price <= 5:
        price_score = 100
    elif price <= 15:
        price_score = 60
    elif price <= 30:
        price_score = 20
    else:
        price_score = 0

    return round(0.5 * review_score + 0.3 * dev_score + 0.2 * price_score, 1)

# ── STEAM API ─────────────────────────────────────────────────────────────────

def extract_appid(url: str) -> str | None:
    m = re.search(r'/app/(\d+)', url.strip())
    return m.group(1) if m else None

def steam_get(url: str, retries=3) -> dict | None:
    for attempt in range(retries):
        try:
            r = requests.get(url, timeout=15, headers={"User-Agent": "Mozilla/5.0"})
            if r.status_code == 429:
                print("  Rate limited, sleeping 30s…")
                time.sleep(30)
                continue
            if r.status_code != 200:
                return None
            return r.json()
        except Exception as e:
            if attempt == retries - 1:
                print(f"  Failed: {e}")
            time.sleep(2)
    return None

def fetch_app_details(appid: str) -> dict | None:
    data = steam_get(f"https://store.steampowered.com/api/appdetails?appids={appid}&l=en")
    if not data or str(appid) not in data:
        return None
    entry = data[str(appid)]
    if not entry.get("success"):
        return None
    return entry.get("data", {})

def fetch_reviews(appid: str) -> int:
    data = steam_get(
        f"https://store.steampowered.com/appreviews/{appid}"
        f"?json=1&language=all&purchase_type=all&num_per_page=0"
    )
    if not data:
        return 0
    return data.get("query_summary", {}).get("total_reviews", 0)

def fetch_dev_game_count(developer: str) -> int:
    if not developer:
        return 1
    data = steam_get(
        f"https://store.steampowered.com/search/results/"
        f"?developer={requests.utils.quote(developer)}&json=1"
    )
    if not data:
        return 1
    return max(1, data.get("total_count", 1))

def parse_price(app_data: dict) -> float:
    if app_data.get("is_free"):
        return 0.0
    po = app_data.get("price_overview", {})
    return po.get("final", 0) / 100.0 if po else 0.0

def parse_release_date(app_data: dict) -> str:
    return app_data.get("release_date", {}).get("date", "")

# ── DISCORD ───────────────────────────────────────────────────────────────────

def discord_post(content: str):
    requests.post(DISCORD_WEBHOOK, json={"content": content}, timeout=10)
    time.sleep(0.5)

def discord_file(file_bytes: bytes, filename: str, content: str = ""):
    requests.post(
        DISCORD_WEBHOOK,
        data={"content": content},
        files={"file": (filename, file_bytes,
               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        timeout=60,
    )

# ── EXCEL ─────────────────────────────────────────────────────────────────────

def build_excel(rows: list[dict]) -> bytes:
    df = pd.DataFrame(rows)[[
        "score", "name", "developer", "genres",
        "release_date", "price_usd", "reviews", "dev_games", "url"
    ]]
    df.index = range(1, len(df) + 1)
    df.index.name = "Rank"
    df.rename(columns={
        "score":        "Score",
        "name":         "Game",
        "developer":    "Developer",
        "genres":       "Genres",
        "release_date": "Release Date",
        "price_usd":    "Price $",
        "reviews":      "Reviews",
        "dev_games":    "Dev's Games",
        "url":          "Steam URL",
    }, inplace=True)

    buf = io.BytesIO()
    df.to_excel(buf, sheet_name="Ranked", index=True)
    buf.seek(0)

    wb = openpyxl.load_workbook(buf)
    ws = wb.active

    # Header
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28

    # Score column color + row font
    score_col = next((c.column for c in ws[1] if c.value == "Score"), None)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = Font(name="Arial", size=9)
            cell.alignment = Alignment(vertical="center")
        if score_col:
            s = row[score_col - 1].value or 0
            if s >= 70:
                row[score_col - 1].fill = PatternFill("solid", fgColor="C6EFCE")
                row[score_col - 1].font = Font(name="Arial", size=9, bold=True, color="276221")
            elif s >= 40:
                row[score_col - 1].fill = PatternFill("solid", fgColor="FFEB9C")
                row[score_col - 1].font = Font(name="Arial", size=9, color="9C5700")
            else:
                row[score_col - 1].fill = PatternFill("solid", fgColor="FFC7CE")
                row[score_col - 1].font = Font(name="Arial", size=9, color="9C0006")

    # Column widths
    for col, w in zip("ABCDEFGHIJ", [7, 9, 34, 24, 20, 12, 9, 9, 9, 48]):
        ws.column_dimensions[col].width = w

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

# ── DISCORD REPORT ────────────────────────────────────────────────────────────

def send_report(rows: list[dict]):
    now = datetime.datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")
    top = rows[:TOP_N]
    zero_reviews = sum(1 for r in rows if r["reviews"] == 0)
    solo_devs    = sum(1 for r in rows if r["dev_games"] == 1)

    discord_post(
        f"## 🎮 Steam Dead Game Report — {now}\n"
        f"**{len(rows)}** games · "
        f"**{zero_reviews}** with 0 reviews · "
        f"**{solo_devs}** solo-dev games\n"
        f"Top **{TOP_N}** ranked by: reviews (50%) + solo dev (30%) + cheap price (20%)\n"
        f"──────────────────────────────────"
    )
    time.sleep(1)

    chunk = []
    for i, g in enumerate(top, 1):
        price_str = "Free" if g["price_usd"] == 0 else f"${g['price_usd']:.2f}"
        medal = "🟢" if g["score"] >= 70 else "🟡" if g["score"] >= 40 else "🔴"
        chunk.append(
            f"{medal} **#{i} {g['name']}** — Score `{g['score']}`\n"
            f"   {price_str} · {g['reviews']} reviews · "
            f"{g['dev_games']} game(s) by dev · {g['release_date']}\n"
            f"   <{g['url']}>"
        )
        if len(chunk) == 10 or i == len(top):
            discord_post("\n".join(chunk))
            chunk = []
            time.sleep(0.5)

    date_str = datetime.datetime.utcnow().strftime("%Y%m%d")
    discord_file(
        build_excel(rows),
        filename=f"steam_dead_{date_str}.xlsx",
        content=f"📊 Full ranked list ({len(rows)} games)",
    )

# ── MAIN ──────────────────────────────────────────────────────────────────────

def main():
    print(f"=== Steam Dead Ranker {datetime.datetime.utcnow()} ===")

    with open(LINKS_FILE) as f:
        appids = [extract_appid(l) for l in f if l.strip()]
    appids = [a for a in appids if a]
    print(f"{len(appids)} links loaded")

    progress = {}
    if os.path.exists(PROGRESS_FILE):
        with open(PROGRESS_FILE) as f:
            progress = json.load(f)
        print(f"Resuming — {len(progress)} already done")

    processed = 0
    for i, appid in enumerate(appids):
        if appid in progress or processed >= MAX_PER_RUN:
            continue

        print(f"[{i+1}/{len(appids)}] {appid}", end=" — ")

        app_data = fetch_app_details(appid)
        time.sleep(BATCH_DELAY_S)

        if not app_data:
            progress[appid] = {"appid": appid, "status": "not_found"}
            print("not found")
            processed += 1
            continue

        developer = (app_data.get("developers") or [""])[0]
        reviews   = fetch_reviews(appid);    time.sleep(BATCH_DELAY_S)
        dev_games = fetch_dev_game_count(developer); time.sleep(BATCH_DELAY_S)
        price     = parse_price(app_data)

        rec = {
            "appid":        appid,
            "status":       "ok",
            "name":         app_data.get("name", ""),
            "developer":    developer,
            "genres":       ", ".join(g["description"] for g in app_data.get("genres", [])),
            "release_date": parse_release_date(app_data),
            "price_usd":    price,
            "reviews":      reviews,
            "dev_games":    dev_games,
            "url":          f"https://store.steampowered.com/app/{appid}",
            "score":        score_game(reviews, dev_games, price),
        }
        progress[appid] = rec
        processed += 1
        print(f"{rec['name'][:45]!r} → {rec['score']}")

        if processed % 50 == 0:
            with open(PROGRESS_FILE, "w") as f:
                json.dump(progress, f)
            print(f"  ✓ saved ({processed} this run)")

    with open(PROGRESS_FILE, "w") as f:
        json.dump(progress, f)

    rows = [v for v in progress.values() if v.get("status") == "ok"]
    rows.sort(key=lambda x: x["score"], reverse=True)
    print(f"\n{len(rows)} valid games. Sending to Discord…")
    send_report(rows)
    print("✓ Done.")

if __name__ == "__main__":
    main()
